"""Layer 2 – PPTX rendering: PresentationContent + the fixed MODY template -> [Product_Name]_MODY_GTM_Deck.pptx.

Responsibilities, and nothing more: open an in-memory copy of the template (the file on disk is never written),
put each mapped field into its named shape, insert the approved hero image without distortion, drop optional groups
whose content is empty, keep every style / layout / slide order of the template, save a new file, and report.
The mapping lives in template_map.py; the text comes from content.py. No layout decision is made here.
"""
import copy, io, os, re, shutil
from validate import ROOT
from presentation import config as cfg
from presentation.template_map import MAP, GROUP_DECOR, SLIDE_COUNT, PLACEHOLDERS, MAPPED
from presentation.content import compose_slots, slot_value, resolve_image, HEB

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
BULLET_INDENT = 171450                                     # EMU (0.19"): hanging indent for bulleted paragraphs


def deck_filename(content):
    name = str(((content.get("presentation") or {}).get("cover") or {}).get("product_name") or "")
    slug = re.sub(r"_+", "_", re.sub(r"[^A-Za-z0-9_-]+", "", re.sub(r"\s+", "_", name.strip()))).strip("_-")
    if not slug:                                             # a Hebrew product name has no ASCII form: use the product id
        slug = str((content.get("meta") or {}).get("product") or "product")
    return slug + cfg.DECK_SUFFIX


# ---------------------------------------------------------------------------------------------------- text
def _first_props(tf):
    """(pPr, rPr) copies of the template paragraph, so every inserted run keeps the designed style."""
    p = tf.paragraphs[0]._p
    pPr = p.find(f"{{{A}}}pPr")
    r = p.find(f"{{{A}}}r")
    rPr = r.find(f"{{{A}}}rPr") if r is not None else p.find(f"{{{A}}}endParaRPr")
    return (copy.deepcopy(pPr) if pPr is not None else None), (copy.deepcopy(rPr) if rPr is not None else None)


def fill_text(shape, items, bullets=False):
    """items: list of str or (bold, rest) pairs -> one paragraph each; [] leaves the shape blank (no filler)."""
    tf = shape.text_frame
    pPr, rPr = _first_props(tf)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]._p
    for child in list(p0):
        if child.tag != f"{{{A}}}pPr":
            p0.remove(child)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pe = p._p
        old = pe.find(f"{{{A}}}pPr")
        if old is not None:
            pe.remove(old)
        ppr = copy.deepcopy(pPr) if pPr is not None else pe.makeelement(f"{{{A}}}pPr", {})
        pe.insert(0, ppr)
        text = "".join(item) if isinstance(item, (tuple, list)) else str(item)
        if HEB.search(text):
            ppr.set("rtl", "1")                            # Hebrew punctuation and mixed Latin names render correctly
        runs = list(item) if isinstance(item, (tuple, list)) else [str(item)]
        if bullets:                                          # the template writes its bullets as literal "• " text
            runs[0] = "• " + str(runs[0])
        for j, chunk in enumerate(runs):
            r = p.add_run()
            r.text = chunk
            if rPr is not None:
                rp = copy.deepcopy(rPr)
                r._r.insert(0, rp)
                if isinstance(item, (tuple, list)) and j == 0 and len(runs) > 1:
                    rp.set("b", "1")                       # inline label ("סטטוס:", the benefit's fact) in bold
    if not items:
        tf.paragraphs[0].add_run().text = ""


def font_pt(shape):
    _, rPr = _first_props(shape.text_frame)
    sz = rPr.get("sz") if rPr is not None else None
    return int(sz) / 100 if sz else 12.0


# a font with Latin + Hebrew glyphs for measuring line breaks (the template's Noto Sans when installed, else a system
# font of similar width); without any, the glyph-average estimate below is used
FONT_CANDIDATES = [os.environ.get("MODY_MEASURE_FONT"), "/Library/Fonts/NotoSans-Regular.ttf", os.path.expanduser("~/Library/Fonts/NotoSans-Regular.ttf"),
                   "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
                   "/System/Library/Fonts/Supplemental/Arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"]
LINE_HEIGHT = 1.2                                          # em, PowerPoint's single spacing for these fonts
_FONT_CACHE = {}


def _measure_font(pt):
    """PIL font at 4× the point size (widths /4 = points), or None when no usable font file exists."""
    key = round(pt * 2) / 2
    if key in _FONT_CACHE:
        return _FONT_CACHE[key]
    font = None
    try:
        from PIL import ImageFont
        path = next((f for f in FONT_CANDIDATES if f and os.path.exists(f)), None)
        font = ImageFont.truetype(path, int(round(key * 4))) if path else None
    except Exception:
        font = None
    _FONT_CACHE[key] = font
    return font


def _insets_pt(shape):
    bp = shape.text_frame._txBody.find(f"{{{A}}}bodyPr")
    l = int(bp.get("lIns", 91440)) if bp is not None else 91440
    r = int(bp.get("rIns", 91440)) if bp is not None else 91440
    t = int(bp.get("tIns", 45720)) if bp is not None else 45720
    b = int(bp.get("bIns", 45720)) if bp is not None else 45720
    return l / 12700, r / 12700, t / 12700, b / 12700


def lines_needed(shape, items, pt):
    """Wrapped line count of the shape's paragraphs at `pt`: measured with a real font when available, else the
    glyph-average estimate (≈ 0.5 em per glyph)."""
    li, ri, _, _ = _insets_pt(shape)
    w_pt = shape.width / 12700 - li - ri
    texts = [p.text_frame.text for p in ()] or ["".join(r.text for r in para.runs) for para in shape.text_frame.paragraphs]
    if not any(t.strip() for t in texts):
        texts = ["".join(i) if isinstance(i, (tuple, list)) else str(i) for i in items]
    font = _measure_font(pt)
    total = 0
    for text in texts:
        if font is None:
            per_line = max(1, int(w_pt / (pt * 0.5)))
            total += max(1, -(-len(text) // per_line))
            continue
        words, cur, n = text.split(), "", 1
        for wd in words:
            cand = (cur + " " + wd).strip()
            if cur and font.getlength(cand) / 4 > w_pt:
                n += 1; cur = wd
            else:
                cur = cand
        total += n
    return total


def overflow_risk(shape, items, pt=None):
    """(over, lines, capacity) at the shape's font size (or `pt`): wrapped lines × 1.2 em against the box height."""
    pt = pt or font_pt(shape)
    _, _, ti, bi = _insets_pt(shape)
    h_pt = shape.height / 12700 - ti - bi
    lines = lines_needed(shape, items, pt)
    capacity = max(1, int(h_pt / (pt * LINE_HEIGHT) + 0.02))
    return lines > capacity, lines, capacity


def shrink_to_fit(shape, items, floor_ratio=0.6, floor_pt=9.0):
    """The template's text boxes declare shrink-on-overflow (<a:normAutofit/>), which PowerPoint only recomputes when
    the text is edited; here the same shrink is applied deterministically so the saved file already fits everywhere.
    Returns (original_pt, final_pt, still_over): the font size is reduced in 0.5 pt steps until the estimate fits,
    never below max(floor_pt, floor_ratio × original). Layout, style and colour are untouched."""
    from pptx.util import Pt
    pt = font_pt(shape)
    risk, _, _ = overflow_risk(shape, items, pt)
    if not risk:
        return pt, pt, False
    floor = max(floor_pt, round(pt * floor_ratio * 2) / 2)
    size = pt
    while size > floor:
        size = round(size - 0.5, 1)
        risk, _, _ = overflow_risk(shape, items, size)
        if not risk:
            break
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)
    return pt, size, risk


# ---------------------------------------------------------------------------------------------------- image
def fit_image(path, box):
    """(left, top, width, height) of `path` fitted inside `box` with its aspect ratio preserved, centred."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    bl, bt, bw, bh = box
    scale = min(bw / iw, bh / ih)
    w, h = int(iw * scale), int(ih * scale)
    return bl + (bw - w) // 2, bt + (bh - h) // 2, w, h


# ---------------------------------------------------------------------------------------------------- render
def render(content, out_dir, template_path=None, product_path=None, filename=None, notes=None):
    """Returns {"pptx": path | None, "warnings": [...], "errors": [...]}. Never raises."""
    template_path = template_path or cfg.TEMPLATE_PATH
    warnings, errors = [], []
    if not os.path.exists(template_path):
        return {"pptx": None, "warnings": warnings, "errors": [f"{cfg.TEMPLATE_NOT_FOUND}: {template_path}"]}
    try:
        from pptx import Presentation
    except ImportError as e:
        return {"pptx": None, "warnings": warnings, "errors": [f"{cfg.RENDER_FAILED}: python-pptx not installed ({e})"]}
    try:
        with open(template_path, "rb") as fh:
            prs = Presentation(io.BytesIO(fh.read()))      # a copy in memory: the master template is never opened for writing
        slides = list(prs.slides)
        if len(slides) != SLIDE_COUNT:
            raise ValueError(f"template has {len(slides)} slides, expected {SLIDE_COUNT}")
        composed = compose_slots(content, final=True)         # the final deck: open items by label, never raw flag syntax
        by_name = [{sh.name: sh for sh in s.shapes} for s in slides]
        # groups whose bound content is empty disappear as a whole (optional elements are hidden, not filled)
        empty_groups = set()
        for name, spec in MAP.items():
            g = spec.get("group")
            if g and spec["kind"] != "caption":
                val = slot_value(content, spec["source"], composed)
                if val in (None, "", []):
                    empty_groups.add(g)
        removed = set()
        for g in empty_groups:
            for name in [n for n, s in MAP.items() if s.get("group") == g] + GROUP_DECOR.get(g, []):
                spec = MAP.get(name, {})
                shapes = by_name[spec["slide"] - 1] if spec else {}
                shape_name = spec.get("shape", name)
                if shape_name in shapes and name not in removed:
                    el = shapes[shape_name]._element
                    el.getparent().remove(el)
                    removed.add(name)
        image_path = resolve_image(slot_value(content, "presentation.cover.hero_image"), product_path or ROOT)
        image_ok = bool(slot_value(content, "presentation.cover.hero_image")) and os.path.exists(image_path)
        if not image_ok:
            warnings.append(f"{cfg.IMAGE_MISSING}: hero image not inserted (placeholder left blank)")
        for name, spec in MAP.items():
            if name in removed:
                continue
            shapes = by_name[spec["slide"] - 1]
            shape_name = spec.get("shape", name)
            if shape_name not in shapes:
                raise KeyError(f"shape {shape_name!r} ({name}) not found on slide {spec['slide']} of {os.path.basename(template_path)}")
            sh = shapes[shape_name]
            kind = spec["kind"]
            if kind == "caption":
                el = sh._element; el.getparent().remove(el)          # the frame stays; its label never survives
                continue
            if kind == "image":
                if image_ok:
                    box = (sh.left, sh.top, sh.width, sh.height)
                    slides[spec["slide"] - 1].shapes.add_picture(image_path, *fit_image(image_path, box))
                continue
            val = slot_value(content, spec["source"], composed)
            if kind == "text":
                items = [] if val in (None, "") else [l for l in str(val).split("\n") if l.strip()]
                fill_text(sh, items)
            else:
                items = list(val or [])
                fill_text(sh, items, bullets=(kind == "bullets"))
            if items:
                orig, size, still = shrink_to_fit(sh, items)
                if size != orig:
                    warnings.append(f"{cfg.TEXT_FIT}: {name} {orig:g} → {size:g} pt (template autofit applied)")
                if still:
                    risk, lines, cap = overflow_risk(sh, items, size)
                    warnings.append(f"{cfg.TEXT_OVERFLOW_RISK}: {name} needs ~{lines} lines, box fits ~{cap} even at {size:g} pt; shorten the text")
        for slide, note in zip(slides, notes or []):            # speaker notes: canonical text, no design
            if note:
                slide.notes_slide.notes_text_frame.text = str(note)
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, filename or deck_filename(content))
        prs.save(out_path)
        left = leftover_placeholders(out_path)
        if left:
            errors.append(f"{cfg.RENDER_FAILED}: placeholder text left in the deck: {left}")
            os.remove(out_path)
            return {"pptx": None, "warnings": warnings, "errors": errors}
        return {"pptx": out_path, "warnings": warnings, "errors": errors}
    except Exception as e:                                  # rendering never blocks the pipeline
        return {"pptx": None, "warnings": warnings, "errors": [f"{cfg.RENDER_FAILED}: {type(e).__name__}: {e}"]}


BRACKET = re.compile(r"\[(?!חסר:)[^\]]+\]")                   # a template placeholder like "[Owner]"; never a MODY flag


def leftover_placeholders(pptx_path):
    from pptx import Presentation
    prs = Presentation(pptx_path)
    left = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and (i, sh.name) in MAPPED:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t in PLACEHOLDERS or BRACKET.search(t) or t.lower().startswith("lorem"):
                        left.append(f"slide {i} {sh.name}: {t!r}")
    return left
