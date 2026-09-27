"""The PPTX renderer applies the template's own shrink-on-overflow rule deterministically (presentation.render.shrink_to_fit)."""
import io, os, sys, unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
try:
    from pptx import Presentation
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False
from presentation import config as cfg, render as R                      # noqa: E402


@unittest.skipUnless(HAVE_PPTX, "python-pptx not installed")
class TestShrinkToFit(unittest.TestCase):
    def shape(self, slide_no, name):
        with open(cfg.TEMPLATE_PATH, "rb") as fh:
            prs = Presentation(io.BytesIO(fh.read()))
        return next(sh for sh in list(prs.slides)[slide_no - 1].shapes if sh.name == name)

    def test_short_text_keeps_the_template_size(self):
        sh = self.shape(7, "Text 17")
        R.fill_text(sh, ["נקודה אחת", "נקודה שנייה"], bullets=True)
        orig, size, still = R.shrink_to_fit(sh, ["נקודה אחת", "נקודה שנייה"])
        self.assertEqual((orig, size, still), (16.0, 16.0, False))

    def test_long_text_is_reduced_until_it_fits_never_below_the_floor(self):
        sh = self.shape(7, "Text 17")
        items = ["זה Jacqueline 77201 של GESSI. הקולקציה הראשונה של Gessi שנוצרה משורשי במבוק אמיתי."] * 4
        R.fill_text(sh, items, bullets=True)
        orig, size, still = R.shrink_to_fit(sh, items)
        self.assertEqual(orig, 16.0); self.assertLess(size, orig); self.assertGreaterEqual(size, max(9.0, round(orig * 0.6 * 2) / 2))
        self.assertFalse(R.overflow_risk(sh, items, size)[0] if not still else False)
        sizes = {run.font.size.pt for p in sh.text_frame.paragraphs for run in p.runs}
        self.assertEqual(sizes, {size})                                                        # every run, one size
        impossible = ["א" * 400] * 12
        R.fill_text(sh, impossible, bullets=True)
        orig, size, still = R.shrink_to_fit(sh, impossible)
        self.assertEqual(size, max(9.0, round(orig * 0.6 * 2) / 2)); self.assertTrue(still)    # floor reached, still flagged


class TestFinalDeckWording(unittest.TestCase):
    """The .pptx never shows raw flag syntax or file names; deck.md keeps the raw flags for QA."""

    def test_flags_read_as_open_items_by_label(self):
        from presentation.content import finalize, flag_label
        self.assertEqual(flag_label("[חסר: technical.flow_rate_lpm – לבדיקה]"), "ספיקה (ל/דק)")
        self.assertEqual(finalize("[חסר: technical.flow_rate_lpm – לבדיקה] · [חסר: technical.installation_type – לבדיקה]"), "ספיקה (ל/דק) · סוג התקנה – להשלמה")
        self.assertEqual(finalize("תאריך השקה: [חסר: launch_plan.launch_date – לבדיקה]"), "תאריך השקה: להגדרה")
        self.assertEqual(finalize('מק"ט: [חסר: product.model_number – לבדיקה]'), 'מק"ט: מק"ט – להשלמה'.replace('מק"ט: מק"ט', 'מק"ט: מק"ט'))
        self.assertEqual(finalize("אין"), "אין"); self.assertIsNone(finalize(None))

    def test_compose_final_vs_markdown(self):
        from presentation.content import compose_slots
        content = {"meta": {"flags": ["[חסר: technical.flow_rate_lpm – לבדיקה]"]},
                   "presentation": {"cover": {}, "product_overview": {}, "positioning": {}, "target_audience": [],
                                    "brand_messaging": {"key_messages": ["מסר ראשון."]}, "product_story": {},
                                    "seller_cheat_sheet": {"opening": "זה X של Y. טענה ארוכה מאוד.", "private_customer": "תועלת א. תועלת ב.",
                                                           "professional_customer": "יצרן איטלקי. MODY – יבואן.", "remember": ["א", "ב", "ג"], "when_to_recommend": ["פרויקט"]},
                                    "launch_plan": [], "next_step": {"owner": "מנהל", "status": "מאומת", "launch_date": "להגדרה",
                                                                     "open_decisions": "[חסר: technical.flow_rate_lpm – לבדיקה]"}}}
        md, final = compose_slots(content), compose_slots(content, final=True)
        self.assertIn("[חסר: technical.flow_rate_lpm – לבדיקה]", md["next_step"]); self.assertNotIn("[חסר:", final["next_step"]); self.assertIn("ספיקה (ל/דק) – להשלמה", final["next_step"])
        self.assertEqual(final["do_dont"], [("לומר: ", "שלושת מסרי המפתח, כלשונם  |  לא לומר: השוואות, סופרלטיבים וטענות ללא מקור")])
        self.assertEqual(final["cheat_how"], [("פתיח: ", "זה X של Y."), ("ללקוח הפרטי: ", "תועלת א"), ("לאדריכלים: ", "MODY – יבואן.")])
        self.assertEqual(final["audience_sales"], [("למי להמליץ: ", "פרויקט"), ("איך לפתוח: ", "זה X של Y."), ("מה להדגיש: ", "תועלת א")])
        self.assertEqual(final["cheat_remember"], ['"א"', '"ב"', '"ג"'])
        from presentation.content import _short_fact
        self.assertEqual(_short_fact("a · b · c · d " * 6), "a · b ועוד"); self.assertEqual(_short_fact("726 – Warm Bronze"), "726 – Warm Bronze")
        self.assertEqual(_short_fact("x" * 80), "x" * 80)                                    # a scalar fact is never cut
        text = " ".join("".join(x) if isinstance(x, tuple) else str(x) for v in final.values() if v for x in (v if isinstance(v, list) else [v]))
        self.assertNotRegex(text, r"sales\.md|landing\.html|launch\.pptx")


if __name__ == "__main__":
    unittest.main()
