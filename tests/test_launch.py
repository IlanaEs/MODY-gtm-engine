"""Tests for the canonical launch package (launch/ package, SKILL step 6-8): single source of truth, package QA before
rendering, independent renderers, and the deck from the fixed MODY template.
Run from the project root:  python -m unittest tests.test_launch -v
Deck rendering tests need python-pptx (requirements.txt) and are skipped without it; everything else needs PyYAML only.
"""
import base64, copy, hashlib, json, os, shutil, subprocess, sys, tempfile, unittest, yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
sys.path.insert(0, ROOT)
from validate import validate                                            # noqa: E402
from qa import run as qa_run                                             # noqa: E402
from launch import package as P                                          # noqa: E402
from launch import qa as pkgqa                                           # noqa: E402
from launch import render_md, render_product_page, deck, pipeline, extract, content_engine   # noqa: E402
from presentation import config as cfg                                   # noqa: E402
from presentation.template_map import SLIDE_COUNT, MAPPED, MAP           # noqa: E402

try:
    import pptx  # noqa: F401
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False

PID, KITCHEN, TILE = "gessi-77201", "gessi-g60077", "mutina-basrelief-patchwork"
ALL = (PID, KITCHEN, TILE)
MD = ("landing.md", "deck.md", "sales.md")


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


def sha_bytes(b):
    return hashlib.sha256(b).hexdigest()


def html_escape(s):
    import html as H
    return H.escape(str(s))


def write_png(path, w, h, rgb=(200, 150, 100)):
    """A flat test bitmap (no Pillow needed): a valid RGB PNG of the given size. Never a product photo."""
    import struct, zlib
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    with open(path, "wb") as fh:
        fh.write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                 + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


class Fixture:
    """products/<pid>.yaml + a temp copy of out/<pid>/ (inside ROOT so relative paths keep working)."""

    def __init__(self, pid=PID, keep_renders=True):
        self.pid = pid
        self.dir = tempfile.mkdtemp(prefix="mody_launch_", dir=ROOT)
        self.out = os.path.join(self.dir, "out") + os.sep
        shutil.copytree(f"out/{pid}", self.out)
        if not keep_renders:
            for f in os.listdir(self.out):
                if f != P.PACKAGE_FILE:
                    os.remove(self.out + f)
        self.product = f"products/{pid}.yaml"
        self.pkg = P.load_package(self.out)

    def ctx(self):
        return P.load_context(self.product, self.out)

    def save(self):
        P.dump(self.pkg, self.out + P.PACKAGE_FILE)

    def qa(self, pkg=None):
        return pkgqa.run(pkg or self.pkg, self.ctx())

    def run(self, **kw):
        self.save()
        return pipeline.run(self.product, self.out, **kw)

    def close(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class ReadyFixture:
    """The same package with a real (synthetic) product image the pipeline can approve: sources.image points at it,
    sources.image_approved is set, and the image-derived facts are refreshed from the skeleton. Nothing is invented:
    the picture is a flat test bitmap, never a product photo."""

    def __init__(self, fx, approved=True, ext="png", size=(800, 300)):
        self.fx = fx
        self.image = os.path.join(fx.dir, f"hero.{ext}")
        write_png(self.image, *size)
        prod = yaml.safe_load(open(fx.product, encoding="utf-8"))
        prod["sources"]["image"] = self.image
        prod["sources"]["image_approved"] = approved
        self.product = os.path.join(fx.dir, "product.yaml")
        yaml.safe_dump(prod, open(self.product, "w", encoding="utf-8"), allow_unicode=True, sort_keys=False)
        self.ctx = P.load_context(self.product, fx.out)
        sk = P.skeleton(self.ctx)
        self.pkg = copy.deepcopy(fx.pkg)
        self.pkg["product"]["assets"], self.pkg["deck"]["launch_plan"], self.pkg["sources"] = sk["product"]["assets"], sk["deck"]["launch_plan"], sk["sources"]
        self.pkg["meta"] = dict(self.pkg["meta"], **{k: sk["meta"][k] for k in ("publishable", "publish_blockers", "flags", "draft_banner")})

    def save(self):
        P.dump(self.pkg, self.fx.out + P.PACKAGE_FILE)

    def run(self, **kw):
        self.save()
        return pipeline.run(self.product, self.fx.out, **kw)


class Base(unittest.TestCase):
    pid = PID

    def setUp(self):
        self.fx = Fixture(self.pid)
        self.pkg = self.fx.pkg

    def tearDown(self):
        self.fx.close()

    def assertError(self, result, code, field_part=None):
        self.assertEqual(result["status"], "failed", result)
        hits = [e for e in result["errors"] if e["code"] == code and (field_part is None or field_part in e["field"])]
        self.assertTrue(hits, f"no {code} error for {field_part!r}: {result['errors']}")

    def assertClean(self, result):
        self.assertIn(result["status"], ("passed", "passed_with_warnings"), result["errors"])


# ---------------------------------------------------------------- package model
class TestPackageModel(Base):
    def test_skeleton_has_every_section_and_facts(self):
        sk = P.skeleton(self.fx.ctx())
        for s in P.SCHEMA["sections"]:
            self.assertIn(s, sk)
        f = sk["product"]["facts"]
        self.assertEqual(f["name"]["value"], "Jacqueline 77201")
        self.assertEqual(f["price"]["display"], "5,000 ₪"); self.assertEqual(f["price"]["source"], "mody_listing")
        self.assertIsNone(f["availability"]["value"]); self.assertIsNone(f["launch_status"]["value"])
        self.assertEqual([r["field"] for r in sk["product"]["specs"]][:4], ["signature_material", "finish", "features", "dimensions_mm"])
        self.assertEqual(sk["product"]["specs"][3]["flag"], "[חסר: technical.dimensions_mm – לבדיקה]")
        self.assertEqual([a["id"] for a in sk["strategy"]["target_audiences"]], ["A", "B"])        # signature -> A, B
        self.assertEqual([s["type"] for s in sk["deck"]["slides"]], [s["type"] for s in P.SCHEMA["deck_slides"]])
        self.assertEqual([st["stage"] for st in sk["deck"]["launch_plan"]], P.SCHEMA["launch_stages"])
        self.assertTrue(sk["story"]["expected"])
        self.assertIsNone(sk["strategy"]["positioning"]["value"])                                   # derived: left to the Skill

    def test_skeleton_for_other_products(self):
        k = P.skeleton(P.load_context(f"products/{KITCHEN}.yaml", f"out/{KITCHEN}/"))
        self.assertFalse(k["story"]["expected"]); self.assertIsNone(k["product"]["facts"]["price"]["value"])
        self.assertEqual([a["id"] for a in k["strategy"]["target_audiences"]], ["C", "B"])
        self.assertEqual(k["meta"]["draft_banner"], "טיוטה – לא לפרסום. חסר: sources.image, commercial.price_ils")
        t = P.skeleton(P.load_context(f"products/{TILE}.yaml", f"out/{TILE}/"))
        self.assertIsNone(t["product"]["facts"]["model"]["value"]); self.assertTrue(t["story"]["expected"])

    def test_committed_packages_are_valid_but_not_ready_without_an_image(self):
        """No product image file exists in the repository: content QA passes, readiness fails, nothing else."""
        for pid in ALL:
            fx = Fixture(pid)
            try:
                r = fx.qa()
                self.assertEqual(r["status"], "not_ready", r["errors"])
                self.assertEqual({k for k, v in r["checks"].items() if v == "FAIL"}, {"readiness"}, r)
                self.assertTrue(all(e["code"] in pkgqa.READINESS_CODES for e in r["errors"]), r["errors"])
                self.assertEqual(fx.pkg["qa"]["status"], r["status"])                                # verdict lives in the package
            finally:
                fx.close()

    def test_canonical_content_exists_once(self):
        """The landing and deck sections only reference canonical fields; they hold no copy of them."""
        for path in ("landing.hero.headline", "landing.hero.sub", "landing.hero.image"):
            self.assertIsInstance(P.getp(self.pkg, path), str)
        self.assertNotIn("positioning", self.pkg["landing"]); self.assertNotIn("benefits", self.pkg["deck"])
        for s in self.pkg["deck"]["slides"]:
            self.assertEqual(set(s), {"id", "type", "source_fields", "notes"})
        self.assertNotIn("legacy", self.pkg["deck"]); self.assertNotIn("opportunity", self.pkg["strategy"])


# ---------------------------------------------------------------- package QA rules
class TestPackageQA(Base):
    def test_exactly_three_benefits(self):
        self.pkg["benefits"].append(dict(self.pkg["benefits"][0]))
        self.assertError(self.fx.qa(), "BENEFIT_COUNT")
        self.pkg["benefits"] = self.pkg["benefits"][:2]
        self.assertError(self.fx.qa(), "BENEFIT_COUNT")

    def test_benefit_needs_fact_and_source(self):
        del self.pkg["benefits"][1]["fact"]
        self.assertError(self.fx.qa(), "BENEFIT_INCOMPLETE", "benefits.1")
        self.pkg["benefits"][1]["fact"] = "x"; self.pkg["benefits"][1]["ref"] = "claim:JQ-9"
        self.assertError(self.fx.qa(), "UNSUPPORTED_CLAIM", "benefits.1")

    def test_technical_fact_without_source_is_rejected(self):
        self.pkg["product"]["specs"][3].update(value="120 × 40 mm", flag=None)
        r = self.fx.qa()
        self.assertError(r, "FACT_MISMATCH", "product.specs")
        self.assertError(r, "MISSING_PRODUCT_SOURCE", "product.specs.3")

    def test_commercial_fact_needs_mody_source(self):
        fx = Fixture(KITCHEN)
        try:
            fx.pkg["product"]["facts"]["price"] = {"value": 1200, "display": "1,200 ₪", "source": "product_page"}
            r = fx.qa()
            self.assertError(r, "MISSING_MODY_SOURCE", "product.facts.price")
            self.assertError(r, "FACT_MISMATCH", "product.facts")
        finally:
            fx.close()

    def test_detail_row_needs_label_value_source(self):
        self.pkg["landing"]["product_details"].append({"label": "אחריות", "value": "10 שנים", "ref": None, "source": None})
        r = self.fx.qa()
        self.assertError(r, "MISSING_PRODUCT_SOURCE", "landing.product_details")
        self.assertError(r, "INVENTED_NUMBER")

    def test_missing_headline_positioning_audience_or_seller_guidance(self):
        self.pkg["messaging"]["headline"]["value"] = ""
        self.pkg["strategy"]["positioning"]["value"] = None
        self.pkg["strategy"]["target_audiences"] = []
        self.pkg["sales"]["seller_cheat_sheet"]["positioning"] = ""
        r = self.fx.qa()
        for field in ("messaging.headline", "strategy.positioning", "strategy.target_audiences", "sales.seller_cheat_sheet.positioning"):
            self.assertError(r, "MISSING_REQUIRED", field)

    def test_hero_requires_headline_and_image_reference(self):
        self.pkg["landing"]["hero"]["image"] = "something.else"
        self.assertError(self.fx.qa(), "HERO_INCOMPLETE", "landing.hero.image")

    def test_missing_or_unapproved_image_blocks_final_assets(self):
        self.assertEqual(P.SCHEMA["image_missing_file"], "error")
        r = self.fx.qa()                                                                              # declared path, file absent
        self.assertError(r, "MISSING_HERO_IMAGE", "hero_image") if r["status"] == "failed" else self.assertEqual(r["status"], "not_ready")
        self.assertTrue(any(e["code"] == "MISSING_HERO_IMAGE" and "not found" in e["message"] for e in r["errors"]))
        fx = Fixture(KITCHEN)                                                                         # no image declared at all
        try:
            r = fx.qa()
            self.assertEqual(r["status"], "not_ready")
            self.assertTrue(any(e["code"] == "MISSING_HERO_IMAGE" for e in r["errors"]))
        finally:
            fx.close()
        ready = ReadyFixture(self.fx, approved=False)                                                 # file exists, MODY has not approved it
        r = pkgqa.run(ready.pkg, ready.ctx)
        self.assertEqual(r["status"], "not_ready"); self.assertTrue(any(e["code"] == "IMAGE_NOT_APPROVED" for e in r["errors"]))
        ready = ReadyFixture(self.fx, approved=True)
        r = pkgqa.run(ready.pkg, ready.ctx)
        self.assertEqual(r["status"], "passed_with_warnings", r["errors"]); self.assertEqual(r["checks"]["readiness"], "PASS")

    def test_unsupported_image_format(self):
        ready = ReadyFixture(self.fx, approved=True, ext="gif")
        r = pkgqa.run(ready.pkg, ready.ctx)
        self.assertEqual(r["status"], "not_ready"); self.assertTrue(any(e["code"] == "IMAGE_FORMAT" for e in r["errors"]))

    def test_deck_needs_cover_and_seven_slides_in_order(self):
        self.pkg["deck"]["slides"].pop(3)
        self.assertError(self.fx.qa(), "DECK_STRUCTURE", "deck.slides")
        self.pkg["deck"]["slides"] = copy.deepcopy(P.SCHEMA["deck_slides"])
        self.pkg["deck"]["cover"]["source_fields"] = []
        self.assertError(self.fx.qa(), "DECK_STRUCTURE", "deck.cover")

    def test_slide_without_content_is_rejected(self):
        self.pkg["strategy"]["consumer_insight"]["value"] = None
        r = self.fx.qa()
        self.assertError(r, "SLIDE_CONTENT_MISSING", "deck.slides.02")
        self.assertError(r, "MISSING_REQUIRED", "strategy.consumer_insight")

    def test_placeholder_text_rejected(self):
        self.pkg["product"]["summary"]["value"] = "Lorem ipsum dolor sit amet."
        self.assertError(self.fx.qa(), "PLACEHOLDER_TEXT", "product.summary")
        self.pkg["product"]["summary"]["value"] = "[Owner]  |  [Status]"
        self.assertError(self.fx.qa(), "PLACEHOLDER_TEXT", "product.summary")

    def test_unsupported_claim_and_missing_derivation(self):
        self.pkg["strategy"]["value_proposition"]["derived_from"] = ["claim:MB-1"]      # another manufacturer's claim
        self.assertError(self.fx.qa(), "UNSUPPORTED_CLAIM", "strategy.value_proposition")
        self.pkg["strategy"]["value_proposition"]["derived_from"] = []
        self.assertError(self.fx.qa(), "MISSING_DERIVATION", "strategy.value_proposition")

    def test_claim_fragment_must_quote_its_item(self):
        self.pkg["story"]["claims"][0]["text"] = "משפט שלא מופיע בסיפור"
        self.assertError(self.fx.qa(), "CLAIM_TEXT", "story.claims.0")

    def test_voice_numbers_price_flags(self):
        self.pkg["strategy"]["customer_need"]["value"] = "הברז הכי מדהים!"
        self.assertError(self.fx.qa(), "VOICE")
        self.pkg["strategy"]["customer_need"]["value"] = "ברז במחיר 4,900 ₪."
        r = self.fx.qa()
        self.assertError(r, "INVENTED_NUMBER"); self.assertError(r, "PRICE")
        self.pkg["strategy"]["customer_need"]["value"] = "צורך."
        self.pkg["deck"]["next_step"]["open_decisions"] = "אין"
        self.pkg["landing"]["product_details"] = [r_ for r_ in self.pkg["landing"]["product_details"] if "dimensions_mm" not in r_["value"]]
        self.pkg["sales"]["faq"] = [q for q in self.pkg["sales"]["faq"] if "dimensions_mm" not in q["a"]]
        r = self.fx.qa()
        self.assertError(r, "FLAG_MISSING")

    def test_other_line_vocabulary_is_rejected(self):
        fx = Fixture(KITCHEN)
        try:
            fx.pkg["strategy"]["positioning"]["value"] = "מלאכת יד מקו Haute Culture."
            self.assertError(fx.qa(), "LINE_LEAKAGE")
        finally:
            fx.close()

    def test_limits(self):
        self.pkg["messaging"]["key_messages"] += [{"value": "4", "derived_from": ["house:P1"]}] * 2
        self.assertError(self.fx.qa(), "LIMIT", "messaging.key_messages")
        self.pkg["messaging"]["key_messages"] = self.pkg["messaging"]["key_messages"][:3]
        self.pkg["story"]["value"] = " ".join(["מילה"] * 81)
        self.assertError(self.fx.qa(), "LIMIT", "story")

    def test_story_only_with_collection_block(self):
        fx = Fixture(KITCHEN)
        try:
            fx.pkg["story"]["value"] = "סיפור מומצא."; fx.pkg["story"]["derived_from"] = ["brand:identity"]
            self.assertError(fx.qa(), "UNSUPPORTED_CLAIM", "story")
        finally:
            fx.close()
        self.pkg["story"]["value"] = None
        self.assertError(self.fx.qa(), "MISSING_REQUIRED", "story")

    def test_meta_bound_to_validator(self):
        self.pkg["meta"]["publishable"] = False
        self.assertError(self.fx.qa(), "META_MISMATCH", "meta.publishable")


# ---------------------------------------------------------------- renderers
class TestRenderers(Base):
    def test_markdown_assets_are_rendered_from_the_package_deterministically(self):
        for pid in ALL:
            fx = Fixture(pid)
            try:
                ctx = fx.ctx()
                self.assertEqual(render_md.landing_md(fx.pkg), open(f"out/{pid}/landing.md", encoding="utf-8").read(), pid)
                self.assertEqual(render_md.deck_md(fx.pkg), open(f"out/{pid}/deck.md", encoding="utf-8").read(), pid)
                self.assertEqual(render_md.sales_md(fx.pkg), open(f"out/{pid}/sales.md", encoding="utf-8").read(), pid)
                prov = yaml.safe_load(open(f"out/{pid}/provenance.yaml", encoding="utf-8"))
                rendered = render_md.provenance(fx.pkg, ctx)
                self.assertEqual({k: v for k, v in prov.items() if k not in ("qa", "renders")}, rendered)
            finally:
                fx.close()

    def test_rendered_markdown_passes_the_32_checks(self):
        fx = Fixture(PID, keep_renders=False)
        try:
            render_md.render(fx.pkg, fx.ctx(), fx.out)
            r = qa_run(fx.product, fx.out)
            self.assertEqual(r["status"], "PASS", r["fails"]); self.assertEqual(len(r["checks"]), 32)
        finally:
            fx.close()

    def test_landing_html_is_the_product_page_template(self):
        self.assertEqual(render_product_page.PAGE_FILE, "landing.html")
        page, _ = render_product_page.render_html(self.pkg, self.fx.product)
        self.assertIn(html_escape(P.val(self.pkg["strategy"]["positioning"])), page)              # canonical positioning
        self.assertIn(html_escape(P.val(self.pkg["messaging"]["headline"])), page)                # main headline
        self.assertIn(html_escape(P.val(self.pkg["messaging"]["one_liner"])), page)               # supporting copy
        for b in self.pkg["benefits"]:
            self.assertIn(html_escape(b["text"]), page)
        self.assertIn(html_escape(self.pkg["story"]["value"]), page)
        self.assertIn("5,000 ₪", page)
        self.assertNotIn("claim:", page); self.assertNotIn("derived_from", page)                  # source metadata is internal

    def test_landing_and_deck_share_canonical_values(self):
        c = deck.content(self.pkg)["presentation"]
        landing = render_md.landing_md(self.pkg)
        page, _ = render_product_page.render_html(self.pkg, self.fx.product)
        self.assertEqual(c["brand_messaging"]["headline"], P.val(self.pkg["messaging"]["headline"]))
        self.assertIn(c["brand_messaging"]["headline"], landing)
        self.assertIn(html_escape(c["positioning"]["positioning_statement"]), page)                # deck and page: same strategy.positioning
        for b in c["product_story"]["benefits"]:
            self.assertIn(html_escape(b["benefit"]), page)
        self.assertEqual(c["cover"]["launch_statement"], P.val(self.pkg["messaging"]["one_liner"]))
        self.assertEqual([b["benefit"] for b in c["product_story"]["benefits"]], [b["title"].rstrip(".") for b in self.pkg["benefits"]])
        self.assertEqual(c["product_overview"]["price"], "5,000 ₪")
        self.assertEqual(c["seller_cheat_sheet"]["remember"], self.pkg["sales"]["key_talking_points"])
        self.assertEqual(len(c["product_story"]["benefits"]), 3)


# ---------------------------------------------------------------- pipeline: QA gate, independence, failure isolation
class TestPipeline(Base):
    def test_not_ready_run_renders_drafts_and_withholds_final_assets(self):
        fx = Fixture(PID, keep_renders=False)
        try:
            for stale in ("landing.html", "launch.pptx"):
                open(fx.out + stale, "w").write("stale")
            r = pipeline.run(fx.product, fx.out)
            self.assertEqual(r["package"], "generated"); self.assertEqual(r["markdown"], "generated"); self.assertEqual(r["markdown_qa"], "PASS")
            self.assertEqual((r["landing"], r["deck"]), ("not_ready", "not_ready"))
            self.assertEqual(r["readiness"]["status"], "not_ready"); self.assertTrue(r["readiness"]["blockers"])
            for f in MD + ("provenance.yaml", pipeline.REPORT, P.PACKAGE_FILE):
                self.assertTrue(os.path.exists(fx.out + f), f)
            for f in ("landing.html", "launch.pptx"):
                self.assertFalse(os.path.exists(fx.out + f), f"stale {f} must not pose as a complete asset")
            self.assertEqual(P.load_package(fx.out)["qa"]["status"], "not_ready")
        finally:
            fx.close()

    def test_ready_run_generates_every_asset_from_one_package_and_one_image(self):
        fx = Fixture(PID, keep_renders=False)
        try:
            ready = ReadyFixture(fx)
            r = ready.run()
            self.assertEqual(r["package"], "generated"); self.assertEqual(r["readiness"]["status"], "ready")
            self.assertEqual((r["markdown"], r["landing"], r["markdown_qa"]), ("generated", "generated", "PASS"), r["errors"])
            page = open(fx.out + "landing.html", encoding="utf-8").read()
            self.assertIn("data:image/png;base64," + base64.b64encode(open(ready.image, "rb").read()).decode()[:60], page)   # the approved photo, embedded
            self.assertIn("#e3202a", page)                                                            # the MODY Product Page Template
            self.assertIn(f"![GESSI Jacqueline 77201]({ready.image})", open(fx.out + "landing.md", encoding="utf-8").read())
            deck_md = open(fx.out + "deck.md", encoding="utf-8").read()
            self.assertEqual(deck_md.count("\n---\n"), 8)                                             # cover + 7 slides + footer rule
            plain = deck_md.replace("**", "")                                                          # markdown emphasis is formatting
            if HAVE_PPTX:
                from pptx import Presentation
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                prs = Presentation(fx.out + "launch.pptx")
                self.assertEqual(len(prs.slides), 8)
                pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and not sh.name.startswith("fixed.")]
                self.assertEqual(len(pics), 2)                                                        # cover + story, same approved file
                self.assertEqual({sha_bytes(p.image.blob) for p in pics}, {sha(ready.image)})
                # parity: every text the deck renderer put on a slide is in deck.md, and the notes are the package notes
                for i, s in enumerate(prs.slides):
                    for sh in s.shapes:
                        if sh.has_text_frame and (i + 1, sh.name) in MAPPED:
                            for p in sh.text_frame.paragraphs:
                                txt = "".join(rn.text for rn in p.runs).strip().lstrip("• ").strip()   # the template's literal bullet
                                if txt and not txt.startswith("house "):
                                    for piece in txt.split(": ", 1):                                  # "label: value" is a table row in markdown
                                        self.assertIn(piece[:40], plain, f"slide {i + 1} {sh.name}: {piece[:40]!r}")
                    self.assertIn(s.notes_slide.notes_text_frame.text.strip(), plain)
        finally:
            fx.close()

    def test_missing_package_is_reported(self):
        fx = Fixture(PID, keep_renders=False)
        try:
            os.remove(fx.out + P.PACKAGE_FILE)
            r = pipeline.run(fx.product, fx.out)
            self.assertEqual(r["package"], "missing"); self.assertTrue(any("LAUNCH_PACKAGE_MISSING" in e for e in r["errors"]))
        finally:
            fx.close()

    def test_deck_failure_keeps_package_and_landing(self):
        ready = ReadyFixture(self.fx)
        r = ready.run(template_path=os.path.join(self.fx.dir, "nope.pptx"))
        self.assertEqual(r["deck"], "render_failed"); self.assertEqual(r["landing"], "generated"); self.assertEqual(r["markdown"], "generated")
        self.assertTrue(any(cfg.TEMPLATE_NOT_FOUND in e for e in r["errors"]))
        pkg = P.load_package(self.fx.out)
        self.assertEqual(pkg["messaging"]["headline"]["value"], self.pkg["messaging"]["headline"]["value"])
        self.assertTrue(os.path.exists(self.fx.out + "landing.html"))
        self.assertFalse(os.path.exists(self.fx.out + "launch.pptx"))

    def test_landing_failure_keeps_package_and_deck(self):
        ready = ReadyFixture(self.fx)
        orig = pipeline.render_product_page.render
        pipeline.render_product_page.render = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
        try:
            r = ready.run()
        finally:
            pipeline.render_product_page.render = orig
        self.assertEqual(r["landing"], "render_failed"); self.assertTrue(any("RENDER_FAILED" in e and "boom" in e for e in r["errors"]))
        self.assertEqual(r["markdown"], "generated"); self.assertEqual(r["package"], "generated")
        if HAVE_PPTX:
            self.assertEqual(r["deck"], "generated")

    def test_run_all_reports_package_landing_and_deck(self):
        r = subprocess.run([sys.executable, "run_all.py"], capture_output=True, text=True, cwd=ROOT)
        head = r.stdout.splitlines()[0]
        for col in ("package", "landing", "deck"):
            self.assertIn(col, head)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        for pid in ALL:
            row = next(l for l in r.stdout.splitlines() if l.startswith(pid))
            self.assertIn("not_ready", row); self.assertIn("PASS", row)                                 # no image file in the repo

    def test_run_all_without_launch_package_still_runs(self):
        d = tempfile.mkdtemp(prefix="mody_launch_")
        try:
            for sub in ("brand", "products", "schema", "out"):
                shutil.copytree(os.path.join(ROOT, sub), os.path.join(d, sub))
            for f in ("validate.py", "qa.py", "run_all.py", "SKILL.md"):
                shutil.copy(os.path.join(ROOT, f), d)
            r = subprocess.run([sys.executable, "run_all.py"], cwd=d, capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr); self.assertNotIn("Traceback", r.stderr); self.assertIn("n/a", r.stdout)
        finally:
            shutil.rmtree(d, ignore_errors=True)


# ---------------------------------------------------------------- product page renderer (fixed template)
class TestProductPage(Base):
    def test_template_is_fixed_and_every_slot_is_filled(self):
        t_sha = sha(render_product_page.TEMPLATE_PATH)
        page, w = render_product_page.render_html(self.pkg, self.fx.product)
        self.assertEqual(sha(render_product_page.TEMPLATE_PATH), t_sha)
        self.assertNotIn("{{", page)
        for token in ("#0a0a0a", "#f2efe9", "#bdb8b0", "#2a2a2a", "#e3202a", "Cormorant+Garamond", "Heebo", "Oswald", "Montserrat", 'dir="rtl"'):
            self.assertIn(token, page)                                                            # the design system, untouched
        self.assertEqual(page.count('class="cta"'), 1)                                            # exactly one red CTA
        self.assertEqual(page.count("<article>"), 3)                                              # 3 advantages
        self.assertEqual(page.count('<h4>'), 4)                                                   # 4 cheat-sheet columns
        self.assertEqual(page.count('<a href="#s0'), 7)                                           # steps 01–07
        self.assertIn("בדיקת זמינות ומחיר", page)
        self.assertIn(html_escape(self.pkg["benefits"][0]["text"]), page)
        self.assertIn(html_escape(P.val(self.pkg["messaging"]["headline"])), page)
        self.assertNotIn("[חסר:", page)                                                           # unknown rows are left out, never guessed
        self.assertNotIn("claim:", page)
        self.assertTrue(any(w_.startswith("images.hero") for w_ in w))                            # no approved photo: reported, not invented

    def test_fields_come_from_the_canonical_package(self):
        d, w = render_product_page.fields(self.pkg, self.fx.product)
        self.assertEqual(d["brand"], "GESSI"); self.assertEqual(d["model"], "77201")
        self.assertEqual(d["pillars"]["core_value"], P.val(self.pkg["strategy"]["value_proposition"]))
        self.assertEqual(d["pillars"]["positioning"], P.val(self.pkg["strategy"]["positioning"]))
        self.assertEqual([a["title"] for a in d["advantages"]], [b["title"].rstrip(".") for b in self.pkg["benefits"]])
        self.assertEqual(d["when_to_recommend"], self.pkg["sales"]["when_to_recommend"][:4])
        self.assertEqual(d["sales_cheat_sheet"]["remember"], self.pkg["sales"]["key_talking_points"])
        self.assertLessEqual(len(d["technical_data"]), 3)
        self.assertTrue(all(row[1] and "[חסר:" not in str(row[1]) for row in d["product_details"]))
        self.assertIsNone(d["images"]["hero"])                                                    # file absent -> nothing embedded

    def test_ready_package_embeds_the_approved_photo(self):
        ready = ReadyFixture(self.fx)
        r = render_product_page.render(ready.pkg, self.fx.out, ready.product)
        self.assertIsNotNone(r["html"], r["errors"])
        page = open(r["html"], encoding="utf-8").read()
        self.assertIn("data:image/png;base64,", page)
        self.assertFalse(any(w_.startswith("images.hero") for w_ in r["warnings"]))
        self.assertEqual(page.count("<button"), 0)                                                # one photo: no thumbnail strip (4 slots only with alternates)
        self.assertTrue(r["html"].endswith("landing.html"))
        bad = os.path.join(self.fx.dir, "bad.html"); open(bad, "w").write("<p>{{unknown_slot}}</p>")
        r = render_product_page.render(ready.pkg, self.fx.out, ready.product, template_path=bad)
        self.assertIsNone(r["html"]); self.assertTrue(any("PRODUCT_PAGE_RENDER_FAILED" in e_ for e_ in r["errors"]))


class TestIntake(unittest.TestCase):
    def test_three_inputs_become_a_product_file_without_facts(self):
        from launch.__main__ import main
        d = tempfile.mkdtemp(prefix="mody_launch_", dir=os.path.join(ROOT, "products"))
        pid = os.path.basename(d) + "-p"
        try:
            rc = main(["intake", pid, "--name", "Test 1", "--url", "https://www.gessi.com/en/bath/test", "--image", "assets/test.png", "--category", "basin_mixer"])
            self.assertEqual(rc, 0)
            path = os.path.join(ROOT, "products", pid + ".yaml")
            doc = yaml.safe_load(open(path, encoding="utf-8"))
            self.assertEqual(doc["product"]["brand"], "gessi")                                    # from the supplier host
            self.assertEqual(doc["sources"]["product_page"], "https://www.gessi.com/en/bath/test")
            self.assertFalse(doc["sources"]["image_approved"]); self.assertEqual(doc["technical"], {})   # no fact invented
            v = validate(path)
            self.assertNotEqual(v["status"], "BLOCKED", v["errors"]) if v["status"] != "BLOCKED" else self.assertTrue(any("specs" in e for e in v["errors"]))
            os.remove(path)
        finally:
            shutil.rmtree(d, ignore_errors=True)
            if os.path.exists(os.path.join(ROOT, "products", pid + ".yaml")):
                os.remove(os.path.join(ROOT, "products", pid + ".yaml"))


# ---------------------------------------------------------------- deck renderer (fixed template)
@unittest.skipUnless(HAVE_PPTX, "python-pptx not installed")
class TestDeck(Base):
    def test_official_template_untouched_and_design_preserved(self):
        from pptx import Presentation
        t_sha, t_mtime = sha(cfg.TEMPLATE_PATH), os.path.getmtime(cfg.TEMPLATE_PATH)
        self.assertTrue(cfg.TEMPLATE_PATH.endswith("assets/templates/presentation/MODY_GTM_Product_Launch_Template.pptx"))
        ready = ReadyFixture(self.fx)
        r = deck.render(ready.pkg, self.fx.out, ready.product)
        self.assertIsNotNone(r["pptx"], r["errors"]); self.assertTrue(r["pptx"].endswith("launch.pptx"))
        self.assertEqual((sha(cfg.TEMPLATE_PATH), os.path.getmtime(cfg.TEMPLATE_PATH)), (t_sha, t_mtime))     # never modified
        out, tpl = Presentation(r["pptx"]), Presentation(cfg.TEMPLATE_PATH)
        self.assertEqual(len(out.slides), SLIDE_COUNT)
        for i, (so, st) in enumerate(zip(out.slides, tpl.slides), 1):
            fixed = lambda s: {sh.name: (sh.left, sh.top, sh.width, sh.height, sh.text_frame.text if sh.has_text_frame else None)
                               for sh in s.shapes if (i, sh.name) not in MAPPED and sh.shape_type != 13}
            self.assertEqual(fixed(so), fixed(st), f"slide {i}: a decorative shape changed")          # design preserved
            captions = {(s["slide"], s["shape"]) for s in MAP.values() if s["kind"] == "caption"}
            for sh in st.shapes:                                                                       # every mapped text slot filled
                if (i, sh.name) in MAPPED and sh.has_text_frame and sh.text_frame.text.strip():
                    filled = next((x for x in so.shapes if x.name == sh.name), None)
                    if (i, sh.name) in captions:
                        self.assertIsNone(filled, f"slide {i} {sh.name}: image caption must be removed")
                    else:
                        self.assertNotEqual(filled.text_frame.text.strip(), sh.text_frame.text.strip(), f"slide {i} {sh.name} still holds template text")
        from presentation.render import leftover_placeholders
        self.assertEqual(leftover_placeholders(r["pptx"]), [])
        texts = " ".join(sh.text_frame.text for s in out.slides for sh in s.shapes if sh.has_text_frame)
        self.assertIn(P.val(self.pkg["messaging"]["headline"]), texts)
        self.assertIn(P.val(self.pkg["strategy"]["positioning"]), texts)
        self.assertIn(self.pkg["sales"]["key_talking_points"][0], texts)
        for f in self.pkg["meta"]["flags"]:
            self.assertIn(f, texts)
        self.assertNotIn("[Owner]", texts); self.assertNotIn("Lorem", texts)

    def test_hero_image_in_both_frames_keeps_aspect_ratio(self):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        ready = ReadyFixture(self.fx, size=(800, 300))
        r = deck.render(ready.pkg, self.fx.out, ready.product)
        self.assertIsNotNone(r["pptx"], r["errors"])
        self.assertFalse(any(cfg.IMAGE_MISSING in w for w in r["warnings"]))
        tpl = Presentation(cfg.TEMPLATE_PATH)
        for slide_no in (2, 6):                                                                        # the template's two image frames
            frame = next(sh for sh in list(tpl.slides)[slide_no - 1].shapes if sh.name == "Shape 6")
            pics = [sh for sh in list(Presentation(r["pptx"]).slides)[slide_no - 1].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pics), 1, f"slide {slide_no}"); pic = pics[0]
            self.assertAlmostEqual(pic.width / pic.height, 800 / 300, places=2)                       # never stretched
            self.assertGreaterEqual(pic.left, frame.left); self.assertLessEqual(pic.left + pic.width, frame.left + frame.width + 1)
            self.assertGreaterEqual(pic.top, frame.top); self.assertLessEqual(pic.top + pic.height, frame.top + frame.height + 1)


# ---------------------------------------------------------------- the actual user flow: name + supplier URL + approved image
FIXTURE_URL = "file://" + os.path.join(ROOT, "tests", "fixtures", "supplier_page.html")


class TestExtraction(unittest.TestCase):
    def test_facts_come_from_the_page_with_evidence(self):
        ex = extract.extract(FIXTURE_URL, "Jacqueline 77201")
        f = ex["facts"]
        self.assertEqual(f["material"]["value"], "Bamboo root + brass"); self.assertIn("Material:", f["material"]["evidence"])
        self.assertEqual(f["features"]["value"][:2], ["Low spout", "Swivel spout"])
        self.assertEqual(f["certifications"]["value"], ["CE", "WRAS"]); self.assertEqual(f["warranty"]["value"], "5 years")
        self.assertEqual(f["sku"]["value"], "77201"); self.assertEqual(f["technical_requirements"]["value"], "1–5 bar")
        self.assertEqual(ex["category"]["value"], "basin_mixer"); self.assertEqual(ex["brand"]["value"], "gessi"); self.assertEqual(ex["line"]["value"], "haute_culture")
        self.assertNotIn("price", f)                                                                  # the script tag is never read as a fact
        doc = extract.product_file("x", "Jacqueline 77201", "https://www.gessi.com/x", "/tmp/i.png", ex)
        self.assertEqual(doc["technical"]["finish"], {"value": {"name": "726 Warm Bronze Brushed PVD"}, "source": "product_page", "evidence": "Finish: 726 Warm Bronze Brushed PVD"})
        self.assertEqual(doc["design"]["claims_ref"], ["JQ-1", "JQ-2", "JQ-3", "JQ-4", "JQ-5"])       # the collection's approved claims only
        self.assertIsNone(doc["commercial"]["price_ils"]); self.assertIsNone(doc["design"]["signature_material"])   # nothing inferred

    def test_unknown_page_yields_no_facts(self):
        d = tempfile.mkdtemp(prefix="mody_launch_", dir=ROOT)
        try:
            page = os.path.join(d, "p.html"); open(page, "w", encoding="utf-8").write("<html><title>Something</title><p>No specs here.</p></html>")
            ex = extract.extract(page, "Thing")
            self.assertEqual(ex["facts"], {}); self.assertIsNone(ex["category"]); self.assertIsNone(ex["brand"])
        finally:
            shutil.rmtree(d, ignore_errors=True)


class TestEndToEnd(unittest.TestCase):
    """Only the three GTM inputs. No product YAML, no launch package, no landing or deck content written by hand."""

    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="mody_launch_", dir=ROOT)
        self.image = os.path.join(self.dir, "hero.png"); write_png(self.image, 800, 1000)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_generate_produces_package_and_both_final_assets(self):
        from launch.__main__ import main
        products, out = os.path.join(self.dir, "products"), os.path.join(self.dir, "out")
        rc = main(["generate", "--name", "Jacqueline 77201", "--url", FIXTURE_URL, "--image", self.image, "--engine", "rules",
                   "--products-root", products, "--out-root", out])
        self.assertEqual(rc, 0 if HAVE_PPTX else 1)                                                  # without python-pptx the deck cannot render
        pid = os.listdir(out)[0]; o = os.path.join(out, pid) + os.sep
        self.assertEqual(pid, "gessi-77201")
        for f in (P.PACKAGE_FILE, "landing.html", pipeline.REPORT, "landing.md", "deck.md", "sales.md", "provenance.yaml") + (("launch.pptx",) if HAVE_PPTX else ()):
            self.assertTrue(os.path.exists(o + f), f)
        pkg = P.load_package(o)
        self.assertIn(pkg["qa"]["status"], ("passed", "passed_with_warnings"))
        report = json.load(open(o + pipeline.REPORT, encoding="utf-8"))
        self.assertEqual((report["package"], report["landing"], report["deck"], report["markdown_qa"]),
                         ("generated", "generated", "generated" if HAVE_PPTX else "render_failed", "PASS"))
        self.assertEqual(report["readiness"]["status"], "ready")
        # facts on both assets come from the page, marketing copy from the engine, the same photo on both
        page = open(o + "landing.html", encoding="utf-8").read()
        self.assertIn("#e3202a", page); self.assertIn("Bamboo root + brass", page); self.assertIn(html_escape(P.val(pkg["strategy"]["positioning"])), page)
        self.assertIn("data:image/png;base64," + base64.b64encode(open(self.image, "rb").read()).decode()[:60], page)
        self.assertEqual(qa_run(os.path.join(products, pid + ".yaml"), o)["status"], "PASS")
        if HAVE_PPTX:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            prs = Presentation(o + "launch.pptx")
            self.assertEqual(len(prs.slides), 8)
            pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
            self.assertEqual(len(pics), 2); self.assertEqual({sha_bytes(p.image.blob) for p in pics}, {sha(self.image)})
            texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
            self.assertIn(P.val(pkg["strategy"]["positioning"]), texts); self.assertIn("Bamboo root + brass", texts)
            from presentation.render import leftover_placeholders
            self.assertEqual(leftover_placeholders(o + "launch.pptx"), [])
        # the generated product file carries evidence, never a value without a source
        prod = yaml.safe_load(open(os.path.join(products, pid + ".yaml"), encoding="utf-8"))
        for field, spec in prod["technical"].items():
            self.assertEqual(spec["source"], "product_page"); self.assertTrue(spec["evidence"], field)

    def test_generate_without_image_file_stops_at_not_ready(self):
        from launch.__main__ import main
        products, out = os.path.join(self.dir, "products"), os.path.join(self.dir, "out")
        rc = main(["generate", "--name", "Jacqueline 77201", "--url", FIXTURE_URL, "--image", os.path.join(self.dir, "missing.png"), "--engine", "rules",
                   "--products-root", products, "--out-root", out])
        self.assertEqual(rc, 1)
        o = os.path.join(out, "gessi-77201") + os.sep
        self.assertTrue(os.path.exists(o + P.PACKAGE_FILE)); self.assertTrue(os.path.exists(o + "landing.md"))
        self.assertFalse(os.path.exists(o + "landing.html")); self.assertFalse(os.path.exists(o + "launch.pptx"))
        self.assertEqual(json.load(open(o + "generate_report.json", encoding="utf-8"))["result"], "not_ready")


class TestContentEngine(Base):
    def test_rules_engine_output_passes_package_qa_for_every_product(self):
        for pid in ALL:
            fx = Fixture(pid, keep_renders=False)
            try:
                ctx = fx.ctx()
                pkg, rep = content_engine.generate(ctx, engine="rules")
                self.assertEqual(rep["engine"], "rules")
                r = pkgqa.run(pkg, ctx)
                self.assertNotEqual(r["status"], "failed", (pid, r["errors"]))                        # only readiness (no image) may block
                self.assertEqual(len(pkg["benefits"]), 3)
                for b in pkg["benefits"]:
                    self.assertTrue(b["ref"]); self.assertIn(b["fact"], b["text"] + b["title"] + " ".join(c["text"] for c in b["claims"]))
                if pid == KITCHEN:
                    self.assertIsNone(pkg["story"]["value"])                                          # no collection block: no story
            finally:
                fx.close()

    def test_claude_engine_falls_back_to_rules_when_unavailable(self):
        orig = content_engine.ClaudeEngine.fill
        content_engine.ClaudeEngine.fill = lambda self, pkg, ctx: (_ for _ in ()).throw(RuntimeError("no credentials"))
        try:
            pkg, rep = content_engine.generate(self.fx.ctx(), engine="claude")
        finally:
            content_engine.ClaudeEngine.fill = orig
        self.assertEqual(rep["engine"], "rules"); self.assertIn("no credentials", rep["fallback"])
        self.assertNotEqual(pkgqa.run(pkg, self.fx.ctx())["status"], "failed")


if __name__ == "__main__":
    unittest.main()
